"""Colour-emoji fix for docx generation.
Splits runs so that traffic-light and status emoji sit in their own runs with
a colour-emoji font applied. Cross-platform: 'Segoe UI Emoji' resolves to the
colour emoji font on Windows Word, and Mac Word falls back to 'Apple Color
Emoji' automatically for emoji code points regardless of the specified font.

Usage:
    from docx import Document
    doc = Document('input.docx')
    fix_colour_emoji(doc)
    doc.save('output.docx')
"""
from copy import deepcopy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Traffic-light and common status emoji used in presales docs.
# Add more code points here if needed (kept explicit rather than a huge range
# so we don't inadvertently 'fix' brand icons or intentional custom fonts).
EMOJI_CHARS = frozenset([
    "\U0001F7E2",  # 🟢 GREEN CIRCLE
    "\U0001F7E1",  # 🟡 YELLOW CIRCLE
    "\U0001F534",  # 🔴 RED CIRCLE
    "\U0001F535",  # 🔵 BLUE CIRCLE
    "\U0001F7E0",  # 🟠 ORANGE CIRCLE
    "\U0001F7E3",  # 🟣 PURPLE CIRCLE
    "\U0001F7E4",  # 🟤 BROWN CIRCLE
    "\U000026AB",  # ⚫ BLACK CIRCLE
    "\U000026AA",  # ⚪ WHITE CIRCLE
    "\U00002705",  # ✅ CHECK MARK BUTTON
    "\U0000274C",  # ❌ CROSS MARK
    "\U000026A0",  # ⚠  WARNING SIGN
])

EMOJI_FONT = "Segoe UI Emoji"


def _segments(text):
    """Split text into (segment_text, is_emoji) groups."""
    if not text:
        return []
    segs, cur, cur_is = [], [], text[0] in EMOJI_CHARS
    for ch in text:
        is_e = ch in EMOJI_CHARS
        if is_e != cur_is:
            segs.append(("".join(cur), cur_is))
            cur, cur_is = [ch], is_e
        else:
            cur.append(ch)
    if cur:
        segs.append(("".join(cur), cur_is))
    return segs


def _set_run_font_to_emoji(r_element):
    """Prepare a run for colour-emoji rendering.

    Word will render colour emoji as MONOCHROME (tinted to the effective run
    colour) whenever any colour is resolved at the run level, including
    colour inherited from the paragraph style, character style, or the
    document's Normal style. Simply removing the run's own <w:color> is not
    enough — Normal's colour still cascades in. To break the chain, we set
    <w:color w:val="auto"/> explicitly on the emoji run. 'auto' is a special
    OOXML value meaning "no specific colour", and Word treats emoji under an
    auto colour as native colour-glyph rendering.

    We also override <w:rFonts> on all four axes to a colour-emoji font."""
    rPr = r_element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        r_element.insert(0, rPr)
    # remove any existing rFonts and colour so we can set them freshly
    for existing in rPr.findall(qn("w:rFonts")):
        rPr.remove(existing)
    for existing_color in rPr.findall(qn("w:color")):
        rPr.remove(existing_color)
    # apply emoji font
    rFonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), EMOJI_FONT)
    rPr.insert(0, rFonts)
    # explicit auto colour to override any inherited colour from styles
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "auto")
    rPr.append(color)


def _run_text_element(r_element):
    """Return the first w:t child, or None."""
    return r_element.find(qn("w:t"))


def _split_run_for_emoji(r_element):
    """If this run's text contains emoji chars, split into sibling runs so
    emoji chars sit in a dedicated run with the emoji font.
    Returns True if any change was made."""
    t = _run_text_element(r_element)
    if t is None or not t.text:
        return False
    if not any(ch in EMOJI_CHARS for ch in t.text):
        return False

    segs = _segments(t.text)
    if len(segs) == 1 and not segs[0][1]:
        return False  # no emoji after all

    parent = r_element.getparent()
    idx = list(parent).index(r_element)

    # CAPTURE the original rPr BEFORE any modification so text clones can
    # inherit the un-modified brand font, not the emoji font.
    original_rPr = r_element.find(qn("w:rPr"))
    original_rPr_snapshot = deepcopy(original_rPr) if original_rPr is not None else None

    # first segment: reuse the existing run
    first_text, first_is_emoji = segs[0]
    t.text = first_text
    t.set(qn("xml:space"), "preserve")
    if first_is_emoji:
        _set_run_font_to_emoji(r_element)

    # remaining segments: insert clones after
    for offset, (seg_text, is_emoji) in enumerate(segs[1:], start=1):
        clone = deepcopy(r_element)
        # replace the text
        clone_t = _run_text_element(clone)
        clone_t.text = seg_text
        clone_t.set(qn("xml:space"), "preserve")
        # restore the pre-modification rPr on the clone (deep-clone from the
        # snapshot so subsequent edits don't affect siblings)
        clone_rPr = clone.find(qn("w:rPr"))
        if clone_rPr is not None:
            clone.remove(clone_rPr)
        if original_rPr_snapshot is not None:
            clone.insert(0, deepcopy(original_rPr_snapshot))
        if is_emoji:
            _set_run_font_to_emoji(clone)
        parent.insert(idx + offset, clone)

    return True


def fix_colour_emoji(doc):
    """Walk every run in body, tables, headers, and footers; split emoji-
    containing runs so emoji render in a colour-emoji font."""
    def _walk_paragraph(p):
        # p is a docx Paragraph or a raw <w:p> element
        p_el = p._p if hasattr(p, "_p") else p
        for r in list(p_el.findall(qn("w:r"))):
            _split_run_for_emoji(r)

    # Body paragraphs
    for p in doc.paragraphs:
        _walk_paragraph(p)

    # Tables (walk cells recursively so nested tables are covered too)
    def _walk_table(tbl):
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _walk_paragraph(p)
                for nested in cell.tables:
                    _walk_table(nested)

    for tbl in doc.tables:
        _walk_table(tbl)

    # Headers and footers
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in hf.paragraphs:
                _walk_paragraph(p)
            for tbl in hf.tables:
                _walk_table(tbl)


if __name__ == "__main__":
    import sys
    from docx import Document
    src = sys.argv[1] if len(sys.argv) > 1 else "input.docx"
    dst = sys.argv[2] if len(sys.argv) > 2 else src.replace(".docx", ".fixed.docx")
    d = Document(src)
    fix_colour_emoji(d)
    d.save(dst)
    print(f"Emoji-fixed: {dst}")
