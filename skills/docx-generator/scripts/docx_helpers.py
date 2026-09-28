"""
docx_helpers.py — shared python-docx recipes for the docx-generator skill.

One import gives every generation batch the same brand-correct building blocks,
so no batch re-types (or mistypes) them:

  plugin_root()                    — resolve the plugin root ($CLAUDE_PLUGIN_ROOT, else this file's location)
  load_brand(name, root)           — read skills/brand/brands/<name>/brand.json from the central registry
  brand_cfg(brand)                 — flatten brand.json into the `cfg` dict every recipe takes
  new_document(brand, root)        — create the document: scratch mode (page setup, logo header,
                                     footer) or template mode (brand letterhead, body cleared)
  open_document(path)              — reopen a saved batch to continue it (multi-batch saving)
  add_heading / add_body / add_bullet / add_numbered
  add_branded_table / add_callout / add_page_break
  add_logo_to_header / add_footer_text / set_cell_bg / set_cell_border
  finalize(doc, path, author)      — colour-emoji fix, author metadata, save, reload check

Usage (inside `uv run --with python-docx==1.1.2 python << 'EOF'`):

    import importlib.util, os
    _ROOT = os.environ.get("CLAUDE_PLUGIN_ROOT", "")
    if not _ROOT:
        raise SystemExit("CLAUDE_PLUGIN_ROOT is not set — substitute the plugin root you read SKILL.md from.")
    _spec = importlib.util.spec_from_file_location(
        "docx_helpers", os.path.join(_ROOT, "skills", "docx-generator", "scripts", "docx_helpers.py"))
    dh = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(dh)

    brand = dh.load_brand("presales-handbook", _ROOT)
    cfg = dh.brand_cfg(brand)
    doc, mode = dh.new_document(brand, _ROOT)
    dh.add_heading(doc, "Executive Summary", 1, cfg)
    dh.add_body(doc, "…", cfg)
    doc.save(OUT)                      # end of batch 1

Depends only on python-docx (pinned 1.1.2 in this plugin) and its bundled lxml.
"""

import importlib.util
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DEFAULT_BRAND = "presales-handbook"
_HERE = Path(__file__).resolve().parent


# ---------------------------------------------------------------------------
# plugin root, brand loading, cfg
# ---------------------------------------------------------------------------

def plugin_root(root=None):
    """Explicit root > $CLAUDE_PLUGIN_ROOT > the plugin folder this file lives in."""
    return str(root or os.environ.get("CLAUDE_PLUGIN_ROOT") or _HERE.parents[2])


def resolve(p, root=None):
    """Expand ${CLAUDE_PLUGIN_ROOT} in a brand.json path. None stays None."""
    return p.replace("${CLAUDE_PLUGIN_ROOT}", plugin_root(root)) if p else None


def load_brand(name=None, root=None):
    """Load the central-registry brand.json. Never a copy inside this skill."""
    path = Path(plugin_root(root)) / "skills" / "brand" / "brands" / (name or DEFAULT_BRAND) / "brand.json"
    if not path.exists():
        raise SystemExit(f"Brand not found: {path} — list skills/brand/brands/ and pick an existing folder.")
    return json.loads(path.read_text(encoding="utf-8"))


def brand_cfg(brand):
    """Flatten brand.json into the dict every recipe uses. Word heading sizes come
    from formats.docx.heading*_pt, never font_sizes.heading*_pt (deck scale)."""
    sem = brand["semantic"]
    dx = brand["formats"]["docx"]
    return {
        "accent": sem["accent"],
        "accent_dark": sem.get("accent_dark", sem["accent"]),
        "text": sem["primary_text"],
        "table_header_bg": sem["table_header_bg"],
        "table_header_text": sem["table_header_text"],
        "callout_bg": sem["callout_bg"],
        "callout_text": sem["callout_text"],
        "callout_border": sem.get("callout_border", sem["accent"]),
        "font_heading": brand["fonts"]["heading"],
        "font_body": brand["fonts"]["body"],
        "font_fallback": brand["fonts"].get("fallback", "Arial"),
        "body_pt": brand["font_sizes"]["body_pt"],
        "heading_pt": {1: dx["heading1_pt"], 2: dx["heading2_pt"], 3: dx["heading3_pt"]},
        "footer_text": dx.get("footer_text", ""),
    }


def hex_to_rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _style_run(run, font, size_pt, color_hex, bold=None):
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.color.rgb = hex_to_rgb(color_hex)
    if bold is not None:
        run.font.bold = bold
    return run


# ---------------------------------------------------------------------------
# document creation (scratch vs template mode) and reopening
# ---------------------------------------------------------------------------

def new_document(brand, root=None, logo_key="logo_dark_png"):
    """Return (doc, mode). Template mode when brand.json has formats.docx.template:
    open it, clear the body, keep its geometry, header and footer. Otherwise scratch
    mode: A4/margins from formats.docx, logo in the header, footer text."""
    dx = brand["formats"]["docx"]
    template = resolve(dx.get("template"), root)
    if template:
        if not Path(template).exists():
            raise SystemExit(f"Brand template not found: {template} — stop and ask the user; "
                             "do not silently switch to scratch mode.")
        doc = Document(template)
        body = doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        return doc, "template"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(dx["page_width_cm"])
    sec.page_height = Cm(dx["page_height_cm"])
    sec.top_margin = Cm(dx["margin_top_cm"])
    sec.bottom_margin = Cm(dx["margin_bottom_cm"])
    sec.left_margin = Cm(dx["margin_left_cm"])
    sec.right_margin = Cm(dx["margin_right_cm"])
    logo = resolve(brand.get("assets", {}).get(logo_key), root)
    if logo and Path(logo).exists():
        add_logo_to_header(doc, logo)
    add_footer_text(doc, dx.get("footer_text", ""), brand_cfg(brand))
    return doc, "scratch"


def open_document(path):
    """Reopen a document saved by an earlier batch. Header, footer, page setup and
    all content are preserved; just keep appending and save to the same path."""
    if not Path(path).exists():
        raise SystemExit(f"No earlier batch at {path} — run batch 1 (new_document) first.")
    return Document(str(path))


# ---------------------------------------------------------------------------
# text recipes
# ---------------------------------------------------------------------------

def add_heading(doc, text, level, cfg):
    level = max(1, min(level, 3))
    try:
        para = doc.add_heading("", level=level)
    except KeyError:  # template without Heading styles
        para = doc.add_paragraph()
    run = para.add_run(text)
    color = cfg["accent"] if level <= 2 else cfg["accent_dark"]
    _style_run(run, cfg["font_heading"], cfg["heading_pt"][level], color, bold=True)
    return para


def add_body(doc, text, cfg, bold=False):
    para = doc.add_paragraph()
    _style_run(para.add_run(text), cfg["font_body"], cfg["body_pt"], cfg["text"], bold=bold or None)
    return para


def _list_para(doc, text, cfg, style, prefix):
    try:
        para = doc.add_paragraph(style=style)
    except KeyError:  # template without list styles
        para = doc.add_paragraph()
        text = prefix + text
    _style_run(para.add_run(text), cfg["font_body"], cfg["body_pt"], cfg["text"])
    return para


def add_bullet(doc, text, cfg, level=0):
    return _list_para(doc, text, cfg, "List Bullet" if level == 0 else "List Bullet 2", "• ")


def add_numbered(doc, text, cfg):
    return _list_para(doc, text, cfg, "List Number", "")


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# tables and callouts
# ---------------------------------------------------------------------------

def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex.lstrip("#"))
    tcPr.append(shd)


def set_cell_border(cell, color_hex, size_eighths=12):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size_eighths))  # eighths of a point
        el.set(qn("w:color"), color_hex.lstrip("#"))
        borders.append(el)
    tcPr.append(borders)


def add_branded_table(doc, headers, rows, cfg, header_pt=10):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_bg(cell, cfg["table_header_bg"])
        _style_run(cell.paragraphs[0].add_run(str(h).upper()), cfg["font_heading"], header_pt,
                   cfg["table_header_text"], bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            _style_run(cell.paragraphs[0].add_run(str(val)), cfg["font_body"], cfg["body_pt"], cfg["text"])
    return table


def add_callout(doc, label, text, cfg, size_pt=10):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, cfg["callout_bg"])
    set_cell_border(cell, cfg["callout_border"])
    para = cell.paragraphs[0]
    _style_run(para.add_run(f"{label}: "), cfg["font_heading"], size_pt, cfg["callout_text"], bold=True)
    _style_run(para.add_run(text), cfg["font_body"], size_pt, cfg["callout_text"])
    return table


# ---------------------------------------------------------------------------
# header and footer (scratch mode — new_document calls these)
# ---------------------------------------------------------------------------

def add_logo_to_header(doc, logo_path, width_cm=3.5):
    header = doc.sections[0].header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.add_run().add_picture(str(logo_path), width=Cm(width_cm))


def add_footer_text(doc, text, cfg, size_pt=8, color="7E8890"):
    if not text:
        return
    footer = doc.sections[0].footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(para.add_run(text), cfg["font_body"], size_pt, color)


# ---------------------------------------------------------------------------
# final batch
# ---------------------------------------------------------------------------

def _emoji_module():
    spec = importlib.util.spec_from_file_location("colour_emoji_fix", _HERE / "colour_emoji_fix.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def fix_colour_emoji(doc):
    """Split 🟢🟡🔴✅❌⚠ into their own colour-emoji runs. Run LAST, after all styling."""
    _emoji_module().fix_colour_emoji(doc)


def finalize(doc, path, author):
    """Final batch only: emoji fix, author metadata, save, then reload to prove the file opens."""
    fix_colour_emoji(doc)
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    Document(str(path))  # raises if the saved file is corrupt
    print(f"Saved and reopened OK: {path}")
    return path
