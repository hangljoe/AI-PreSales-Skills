# /// layout
# purpose = "Opening slide for a presentation — company / deck name, subtitle, optional date; also section dividers"
# best_for = "First slide of any deck; section dividers with a major heading (add_section_slide)"
# avoid_when = "Content slides; anything with bullet points or data"
# max_title_chars = 60
# max_subtitle_chars = 120
# instructions = """
# Title slides use tok["title_slide_bg"], section dividers tok["section_slide_bg"]
# (navy #112D4E for presales-handbook, from brand.json formats.pptx) — NOT the
# white content canvas. Text colour comes from tok["title_slide_text"] /
# tok["section_slide_text"] (white on a dark canvas).
# Large centred title; subtitle below in a softer tone.
# Logo: pass logo=bh.logo_path(brand, on_dark=bh.is_dark(bg)) — the white logo on
# navy, the dark logo on a light canvas. Never reference a logo copy in this skill.
# Optional date or account name in small text at bottom-right.
# No bullet points. No body copy. Title is the whole message.
# """
# ///

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def hex_to_rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

def add_title_slide(prs, title, subtitle="", footer="",
                    bg="#112D4E", primary="#FFFFFF", secondary="#E7E4E2",
                    logo=None, logo_width=Inches(2.2)):
    """Defaults are the presales-handbook title canvas (navy, white text).
    For any brand pass bg=tok["title_slide_bg"], primary=tok["title_slide_text"];
    on a light title canvas also pass secondary=tok["text"]."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg_fill = slide.background.fill
    bg_fill.solid()
    bg_fill.fore_color.rgb = hex_to_rgb(bg)

    W, H = prs.slide_width, prs.slide_height

    # Logo — top-left, from the brand registry (see instructions above)
    if logo:
        slide.shapes.add_picture(logo, Inches(0.6), Inches(0.5), width=logo_width)

    # Title — centred, large
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(1.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb(primary)

    # Subtitle
    if subtitle:
        tx2 = slide.shapes.add_textbox(Inches(1.0), Inches(4.2), W - Inches(2.0), Inches(1.0))
        tf2 = tx2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(18)
        r2.font.color.rgb = hex_to_rgb(secondary)

    # Footer
    if footer:
        tx3 = slide.shapes.add_textbox(W - Inches(3.5), H - Inches(0.7), Inches(3.0), Inches(0.4))
        tf3 = tx3.text_frame
        p3 = tf3.paragraphs[0]
        p3.alignment = PP_ALIGN.RIGHT
        r3 = p3.add_run()
        r3.text = footer
        r3.font.size = Pt(10)
        r3.font.color.rgb = hex_to_rgb(secondary)

    return slide

def add_section_slide(prs, title, subtitle="", tok=None, logo=None):
    """Section divider: same layout on tok["section_slide_bg"] (navy by default)."""
    tok = tok or {}
    bg = tok.get("section_slide_bg", "112D4E")
    primary = tok.get("section_slide_text", "FFFFFF")
    # White text means a dark canvas -> soft light subtitle; otherwise body text.
    secondary = "E7E4E2" if primary == "FFFFFF" else tok.get("text", "322F2F")
    return add_title_slide(
        prs, title, subtitle,
        bg=bg, primary=primary, secondary=secondary,
        logo=logo, logo_width=Inches(1.6),
    )
