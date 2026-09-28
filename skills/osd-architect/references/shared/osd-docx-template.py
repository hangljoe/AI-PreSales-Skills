"""
OSD Word generator — branded Opportunity Scoping Document (.docx).

Run with:
    uv run --with python-docx==1.1.2 python <your-completed-script>.py

This is a STARTING POINT / helper library. The osd-architect skill fills in:
  - BRAND (presales-handbook default) and SHAREABLE (False = full internal OSD; True = customer-safe)
  - the account / SC / AE / PS header values
  - every section via the add_* helpers, in osd-structure.md order (handbook chapter 8.2), with
    the SC's product scope in Section 9 and rendered diagram PNGs embedded via add_image.

Slick features built in:
  - Cover page + clickable Table of Contents (headings use built-in styles so the TOC populates)
  - add_image() embeds rendered diagram PNGs (business flow map, As-Is/To-Be flows, IT architecture),
    with a labelled placeholder fallback when the PNG isn't present
  - Confidence legend + inline tags render as real coloured circles (not emoji, which Word
    greys out) via colourize_dots(); add_gaps_page() lists every 🔴 Unknown as an action list
  - SHAREABLE mode: skip internal/commercial sections for the customer-facing copy
  - Fields auto-update on open (so the TOC fills in without a manual refresh in most clients)

For UPDATE-IN-PLACE of an existing OSD, use osd-docx-update.py instead.
"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# === CONFIG (set by the skill) ===
BRAND = "presales-handbook"   # or the folder name of your own brand under skills/brand/brands/
SHAREABLE = False     # True => customer-safe redacted copy (shareable sections only)
# IMPORTANT for SHAREABLE: wrap each INTERNAL section's whole body in
#   if add_heading1("…", internal=True):
#       …body…
# add_heading1(internal=True) returns False in shareable mode, so the guarded body is skipped.
# Guarding only the heading (not the body) will LEAK commercial/MEDDPICC content — don't do that.

# Absolute plugin root — SET BY THE SKILL (osd-architect Step 5a writes the resolved absolute path
# here). It is what makes a custom brand.json and the logo load. ${CLAUDE_PLUGIN_ROOT} is NOT
# expanded inside a Python string, so the env var is only a fallback. If neither is set, the
# presales-handbook default colours apply, the logo is skipped, and a NOTE is printed.
PLUGIN_ROOT = ""
_ROOT = PLUGIN_ROOT or os.environ.get("CLAUDE_PLUGIN_ROOT", "")
if not _ROOT:
    print("NOTE: PLUGIN_ROOT not set — using presales-handbook default brand values, no logo.")

# Anchor for resolving diagram PNGs (#23). Diagrams live under the RUN directory
# (`output/diagrams/…`), NOT wherever the interpreter's CWD happens to be. Anchoring
# to this script's own location makes add_image() work regardless of CWD. Falls back
# to CWD when __file__ is unavailable (e.g. run via a heredoc).
try:
    RUN_DIR = os.path.dirname(os.path.abspath(__file__))
except NameError:
    RUN_DIR = os.getcwd()

# Brand: presales-handbook by default (skills/brand/brands/presales-handbook/brand.json).
# If that file exists it is read; otherwise the hardcoded defaults apply (navy 112D4E,
# yellow FACF39, text 322F2F, Arial). Users can add their own brand folder under
# skills/brand/brands/<name>/ and set BRAND to its folder name.
import json as _json
_BRAND_DIR = os.path.join(_ROOT or ".", "skills", "brand", "brands", BRAND)
_brand = {}
try:
    with open(os.path.join(_BRAND_DIR, "brand.json"), encoding="utf-8") as _fh:
        _brand = _json.load(_fh)
except (OSError, ValueError):
    pass
_c = _brand.get("colors", {}); _s = _brand.get("semantic", {}); _f = _brand.get("fonts", {})
ACCENT_HEX = (_s.get("accent") or _c.get("navy") or _c.get("primary") or "112D4E").lstrip("#").upper()
HIGHLIGHT_HEX = (_s.get("highlight") or _c.get("yellow") or _c.get("secondary") or "FACF39").lstrip("#").upper()
TEXT_HEX = (_s.get("primary_text") or _c.get("text") or "322F2F").lstrip("#").upper()
ACCENT = RGBColor.from_string(ACCENT_HEX); HIGHLIGHT = RGBColor.from_string(HIGHLIGHT_HEX)
TEXT = RGBColor.from_string(TEXT_HEX)
FONT = _f.get("heading") or _f.get("body") or "Arial"
BRAND_LABEL = _brand.get("display_name") or _brand.get("name") or "The PreSales Handbook"
LOGO = os.path.join(_BRAND_DIR, "assets", "logo-dark.png") if _ROOT else ""  # optional — skipped if absent

MUTED = RGBColor(0x7A, 0x7A, 0x7A)
# Confidence-tag colours (semantic, brand-independent) — drive the legend and colourize_dots()
GREEN = RGBColor(0x2E, 0x7D, 0x32); AMBER = RGBColor(0xE0, 0x9B, 0x00); RED = RGBColor(0xC0, 0x39, 0x2B)
DOTS = {"🟢": GREEN, "🟡": AMBER, "🔴": RED}

# === METADATA (replace) ===
ACCOUNT = "[Account Name]"; AREA = "[Product Scope]"; SC = "[SC]"; AE = "[AE]"
PS = "[PS]"; DATE = "[Date]"; VERSION = "1.0"

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21); section.page_height = Cm(29.7)
section.left_margin = section.right_margin = Cm(2.2)
section.top_margin = section.bottom_margin = Cm(1.8)

# Brand the built-in heading styles so the TOC picks them up AND they look on-brand.
for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
    st = doc.styles[name]
    st.font.name = FONT; st.font.size = Pt(size)
    st.font.color.rgb = ACCENT if name == "Heading 1" else TEXT
    st.font.bold = True

# === HELPERS ===
def _logo_header():
    hp = section.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(LOGO):
        hp.add_run().add_picture(LOGO, width=Inches(1.5))

def add_cover():
    if os.path.exists(LOGO):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(LOGO, width=Inches(2.4))
    for _ in range(4): doc.add_paragraph()
    t = doc.add_paragraph(); r = t.add_run("Opportunity Scoping Document")
    r.font.name = FONT; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = ACCENT
    _rule(HIGHLIGHT_HEX)
    s = doc.add_paragraph(); rs = s.add_run(f"{ACCOUNT} — {AREA}")
    rs.font.name = FONT; rs.font.size = Pt(16); rs.font.color.rgb = TEXT
    doc.add_paragraph()
    meta = (("Client", ACCOUNT), ("Solution Consultant", SC), ("Account Executive (AE)", AE),
            ("Professional Services (PS)", PS), ("Version", VERSION), ("Date", DATE),
            ("Classification", "Customer-shareable" if SHAREABLE else "Internal — Confidential"))
    for k, v in meta:
        p = doc.add_paragraph()
        a = p.add_run(f"{k}:  "); a.font.bold = True; a.font.name = FONT; a.font.size = Pt(10); a.font.color.rgb = TEXT
        b = p.add_run(str(v)); b.font.name = FONT; b.font.size = Pt(10); b.font.color.rgb = TEXT
    doc.add_page_break()

def _rule(hex_):
    """A short brand-highlight rule under the cover title (bottom border on an empty paragraph)."""
    p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "1"), ("w:color", hex_)):
        bot.set(qn(k), v)
    bdr.append(bot); pPr.append(bdr)

def add_toc():
    h = doc.add_paragraph(); r = h.add_run("Contents")
    r.font.name = FONT; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = ACCENT
    p = doc.add_paragraph(); run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = 'TOC \\o "1-2" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    tt = OxmlElement("w:t"); tt.text = "Update this field (Ctrl+A, F9) to build the table of contents."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, tt, f3): run._r.append(e)
    doc.add_page_break()

def _update_fields_on_open():
    el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
    doc.settings.element.append(el)


def bake_toc(docx_path):
    """Compute and bake the Table of Contents into the saved .docx via LibreOffice UNO,
    so the Contents page populates immediately on open without requiring a manual
    Ctrl+A / F9 refresh. Falls back silently if LibreOffice is not available."""
    import subprocess, shutil, sys as _sys, time, tempfile, os
    if not shutil.which("soffice"):
        print("[bake_toc] soffice not found — TOC will need Ctrl+A / F9 on first open.", file=_sys.stderr)
        return
    profile = tempfile.mkdtemp(prefix="lo_osd_")
    uno_script = (
        "import sys, time, uno\n"
        "from com.sun.star.beans import PropertyValue\n"
        "def pv(n,v):\n    p=PropertyValue(); p.Name=n; p.Value=v; return p\n"
        "def connect(r=30):\n"
        "    lc=uno.getComponentContext()\n"
        "    res=lc.ServiceManager.createInstanceWithContext(\"com.sun.star.bridge.UnoUrlResolver\",lc)\n"
        "    for _ in range(r):\n"
        "        try: return res.resolve(\"uno:socket,host=127.0.0.1,port=2002;urp;StarOffice.ComponentContext\")\n"
        "        except: time.sleep(0.5)\n"
        "    raise RuntimeError(\"could not connect\")\n"
        "ctx=connect()\n"
        "desktop=ctx.ServiceManager.createInstanceWithContext(\"com.sun.star.frame.Desktop\",ctx)\n"
        "url=\"file://\"+sys.argv[1]\n"
        "doc=desktop.loadComponentFromURL(url,\"_blank\",0,(pv(\"Hidden\",True),))\n"
        "idxs=doc.getDocumentIndexes()\n"
        "[idxs.getByIndex(i).update() for i in range(idxs.getCount())]\n"
        "try: doc.refresh()\nexcept: pass\n"
        "doc.storeToURL(url,(pv(\"FilterName\",\"MS Word 2007 XML\"),pv(\"Overwrite\",True)))\n"
        "doc.close(False)\n"
    )
    script_path = os.path.join(profile, "bake_toc_uno.py")
    with open(script_path, "w") as f: f.write(uno_script)
    so = subprocess.Popen(
        ["soffice", "--headless", "--invisible", "--norestore", "--nologo",
         "--accept=socket,host=127.0.0.1,port=2002;urp;",
         f"-env:UserInstallation=file://{profile}"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        time.sleep(8)
        subprocess.run([_sys.executable, script_path, os.path.abspath(docx_path)],
                       timeout=60, check=True)
    except Exception as e:
        print(f"[bake_toc] failed ({e}) — TOC needs manual Ctrl+A / F9.", file=_sys.stderr)
    finally:
        so.terminate(); so.wait(timeout=10)
        import shutil as _sh2; _sh2.rmtree(profile, ignore_errors=True)

def add_heading1(text, internal=False):
    if internal and SHAREABLE: return False
    doc.add_paragraph(text, style="Heading 1")
    return True

def add_heading2(text):
    doc.add_paragraph(text, style="Heading 2")

def add_para(text, italic=False, muted=False):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(10); r.font.italic = italic
    r.font.color.rgb = MUTED if muted else TEXT
    return p

def add_bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); r = p.add_run(it)
        r.font.name = FONT; r.font.size = Pt(10); r.font.color.rgb = TEXT

def add_banner(text):
    p = doc.add_paragraph(); r = p.add_run("⚠ " + text)
    r.font.name = FONT; r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = MUTED

def add_confidence_legend():
    # Real coloured circles, not emoji: Word tints emoji to monochrome when a font
    # colour is applied, which makes the key render greyscale. Solid circles keep colour.
    p = doc.add_paragraph()
    def _seg(txt, col, dot=False):
        r = p.add_run(txt); r.font.name = FONT; r.font.size = Pt(11 if dot else 9)
        r.font.italic = not dot; r.font.color.rgb = col
    _seg("Confidence:   ", MUTED)
    _seg("●", GREEN, dot=True); _seg("  Confirmed (sourced)      ", MUTED)
    _seg("●", AMBER, dot=True); _seg("  Inferred (reasoned)      ", MUTED)
    _seg("●", RED, dot=True);   _seg("  Unknown (to confirm)", MUTED)


def colourize_dots(document):
    """Convert inline confidence emoji in the finished document into solid coloured
    circles, so the tags render in colour in Word (emoji get tinted to monochrome when
    a font colour is applied). Call once, just before saving."""
    def _fix(paragraph):
        full = "".join(run.text for run in paragraph.runs)
        if not any(d in full for d in DOTS):
            return
        base = paragraph.runs[0]
        nm = base.font.name or FONT; sz = base.font.size or Pt(9)
        bd = base.font.bold; itx = base.font.italic
        try:
            col = base.font.color.rgb if (base.font.color is not None and base.font.color.type is not None) else None
        except Exception:
            col = None
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        for seg in re.split(r"(🟢|🟡|🔴)", full):
            if not seg:
                continue
            is_dot = seg in DOTS
            nr = paragraph.add_run("●" if is_dot else seg)
            nr.font.name = nm; nr.font.bold = bd; nr.font.italic = itx
            nr.font.size = Pt(10) if is_dot else sz
            if is_dot:
                nr.font.color.rgb = DOTS[seg]
            elif col is not None:
                nr.font.color.rgb = col
    for p in document.paragraphs:
        _fix(p)
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fix(p)

def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_); tcPr.append(shd)

def _style_cell(cell, bold=False, header=False):
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = FONT; run.font.size = Pt(9); run.font.bold = bold or header
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if header else TEXT

def add_kv_table(rows):
    t = doc.add_table(rows=len(rows), cols=2); t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        t.rows[i].cells[0].text = str(k); t.rows[i].cells[1].text = str(v)
        _style_cell(t.rows[i].cells[0], bold=True); _style_cell(t.rows[i].cells[1])

def add_grid_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for j, hdr in enumerate(headers):
        t.rows[0].cells[j].text = str(hdr); _shade(t.rows[0].cells[j], ACCENT_HEX)
        _style_cell(t.rows[0].cells[j], header=True)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val); _style_cell(cells[j])

def _resolve_png(png_path):
    """
    Resolve a diagram PNG path CWD-independently (#23). Tries, in order:
      1. the path as given (absolute, or relative to CWD)
      2. relative to this script's RUN_DIR
      3. relative to RUN_DIR/output (so 'diagrams/x.png' finds 'output/diagrams/x.png')
    Returns the first existing path, else None.
    """
    if not png_path:
        return None
    candidates = [png_path,
                  os.path.join(RUN_DIR, png_path),
                  os.path.join(RUN_DIR, "output", png_path)]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None

def add_image(png_path, caption="", source_note="", width_inches=6.3):
    """Embed a rendered diagram PNG; fall back to a labelled placeholder box.

    #23: resolves paths regardless of CWD and prints a LOUD warning (not a silent
    placeholder) when a diagram is missing — so a document is never delivered with
    invisible gaps that look intentional."""
    resolved = _resolve_png(png_path)
    if resolved:
        doc.add_picture(resolved, width=Inches(width_inches))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption: add_para(caption, italic=True, muted=True)
    else:
        print(f"WARNING [add_image]: diagram NOT FOUND — {png_path!r}. "
              f"Searched CWD ({os.getcwd()}) and RUN_DIR ({RUN_DIR}). "
              f"Embedding a placeholder for '{caption or 'Diagram'}'. "
              f"Render the diagram (see diagrams.md) and re-run before delivering.")
        t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
        c = t.rows[0].cells[0]
        c.text = f"[ {caption or 'Diagram'} — to be inserted ]"
        if source_note:
            c.add_paragraph(source_note)
        for para in c.paragraphs:
            for run in para.runs:
                run.font.italic = True; run.font.color.rgb = MUTED

def add_gaps_page(gaps):
    """gaps = [(section, item, owner, target_date), ...] — every 🔴 Unknown."""
    add_heading1("Gaps to Confirm")
    add_para("Every 🔴 Unknown captured in this OSD. Drive these in the next conversation.", muted=True)
    if gaps:
        add_grid_table(["#", "Section", "Open item to confirm", "Owner", "Target"],
                       [[i + 1, s, it, o, d] for i, (s, it, o, d) in enumerate(gaps)])
    else:
        add_para("No open gaps recorded.")

# === ASSEMBLY (the skill replaces the demo content below) ===
_logo_header()
add_cover()
add_toc()

add_heading1("Document History")
add_grid_table(["Version", "Date", "Name", "Description"], [[VERSION, DATE, SC, "Initial draft"]])

if add_heading1("1 — Executive Summary", internal=True):
    add_confidence_legend()
    add_heading2("Customer and Opportunity Information")
    add_kv_table([("Account Name", ACCOUNT), ("Industry", "[…]"), ("Compelling Event", "[…] 🔴")])
    add_banner("Commercial fields (ARR, competition, MEDDPICC) are internal-only.")

add_heading1("5 — Supply Chain / Business Flow Map & As-Is Processes")   # shareable
add_heading2("Supply Chain / Business Flow Map")
add_image(png_path="output/diagrams/sc-map.png", caption="As-Is business flow map",
          source_note="Generated via the diagram skill from account research and discovery.")

add_heading1("6 — IT Overview & Architecture")                           # shareable
add_image(png_path="output/diagrams/to-be-architecture.png", caption="To-Be IT architecture")

add_gaps_page([("Executive Summary", "Confirm compelling event & target go-live", "SC", "[date]")])

colourize_dots(doc)
_update_fields_on_open()

out = f"output/OSD-{ACCOUNT}-{AREA}-v{VERSION}{'-SHAREABLE' if SHAREABLE else ''}.docx".replace(" ", "_")
os.makedirs("output", exist_ok=True)
doc.save(out)
bake_toc(out)
print(f"Saved: {out}")
