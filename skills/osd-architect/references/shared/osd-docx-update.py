"""
OSD update-in-place — revise an existing OSD .docx without regenerating it.

Run with:
    uv run --with python-docx==1.1.2 python <your-completed-update-script>.py

OSDs are living documents (single source of truth, updated after every call). This module
opens the existing .docx and edits it surgically:

  - append_history_row(doc, ...)  → adds a row to the Document History table
  - replace_section(doc, "6 — IT Overview & Architecture", fill_fn) → replaces the body of one H1 section
  - append_to_section(doc, heading, fill_fn) → appends content to the end of a section

The skill reads the existing OSD first (text + tables) to understand current content, decides
what changed, then calls these helpers. Always bump the version and append a history row so the
change log stays honest. Unchanged sections and any manual SC formatting are preserved.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

DOC_PATH = "output/OSD-[Account]-[Product]-v1.0.docx"   # set by the skill (the existing OSD)
NEW_VERSION = "1.1"
# presales-handbook brand defaults (see osd-docx-template.py); match the brand the OSD was built with.
FONT = "Arial"; TEXT = RGBColor(0x32, 0x2F, 0x2F)


def _is_h1(el):
    if el.tag != qn("w:p"):
        return False
    pPr = el.find(qn("w:pPr"))
    if pPr is None:
        return False
    pStyle = pPr.find(qn("w:pStyle"))
    return pStyle is not None and pStyle.get(qn("w:val")) in ("Heading1", "Heading 1")


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def append_history_row(doc, version, date, name, description):
    """Append a row to the Document History table (first table whose header has 'Version')."""
    for t in doc.tables:
        header = " ".join(c.text for c in t.rows[0].cells).lower()
        if "version" in header and ("date" in header or "name" in header):
            cells = t.add_row().cells
            for cell, val in zip(cells, [version, date, name, description]):
                cell.text = str(val)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = FONT; r.font.size = Pt(9); r.font.color.rgb = TEXT
            return True
    raise ValueError("Document History table not found")


def _relocate_after(body, anchor, new_els):
    for e in new_els:
        body.remove(e)
    for e in new_els:
        anchor.addnext(e); anchor = e


def _capture_new(doc, fill_fn):
    """Run fill_fn (which appends blocks at the end) and return the new top-level elements."""
    body = doc.element.body
    before = set(id(c) for c in body.iterchildren())
    fill_fn()
    return body, [c for c in body.iterchildren() if id(c) not in before]


def replace_section(doc, heading_text, fill_fn):
    """Delete the body between a Heading-1 and the next Heading-1, then insert fill_fn's output."""
    body = doc.element.body
    target = next((el for el in body.iterchildren()
                   if _is_h1(el) and _text(el).strip() == heading_text.strip()), None)
    if target is None:
        raise ValueError(f"Section not found: {heading_text}")
    el = target.getnext()
    while el is not None and not _is_h1(el) and el.tag != qn("w:sectPr"):
        nxt = el.getnext(); body.remove(el); el = nxt
    _, new_els = _capture_new(doc, fill_fn)
    _relocate_after(body, target, new_els)


def append_to_section(doc, heading_text, fill_fn):
    """Append fill_fn's output to the END of a section (before the next Heading-1)."""
    body = doc.element.body
    target = next((el for el in body.iterchildren()
                   if _is_h1(el) and _text(el).strip() == heading_text.strip()), None)
    if target is None:
        raise ValueError(f"Section not found: {heading_text}")
    el = target.getnext(); last = target
    while el is not None and not _is_h1(el) and el.tag != qn("w:sectPr"):
        last = el; el = el.getnext()
    _, new_els = _capture_new(doc, fill_fn)
    _relocate_after(body, last, new_els)


# === EXAMPLE USAGE (skill replaces this) ===
if __name__ == "__main__":
    doc = Document(DOC_PATH)

    def new_asis():
        p = doc.add_paragraph(); r = p.add_run(
            "Updated after the 12 June technical call: confirmed SAP ECC + MuleSoft middleware. 🟢")
        r.font.name = FONT; r.font.size = Pt(10); r.font.color.rgb = TEXT

    replace_section(doc, "6 — IT Overview & Architecture", new_asis)
    append_history_row(doc, NEW_VERSION, "[date]", "[SC]", "Updated IT overview after technical call")

    out = DOC_PATH.replace("v1.0", f"v{NEW_VERSION}")
    doc.save(out)
    print(f"Saved: {out}")
