"""
OSD scoping export → Excel for PS estimation.

Run with:
    uv run --with openpyxl==3.1.5 python <your-completed-script>.py

Exports every Modules & Detailed Scoping answer (the ROM / SOW / CF rows) into a clean,
filterable workbook the PS team uses for the estimation / SOW process. One row per requirement,
tagged with what it's needed for (ROM = Rough Order of Magnitude, SOW = Statement of Work,
CF = Complexity Factor) and a confidence tag.

The skill fills ROWS from OSD Sections 6.3 and 9.2. Group is the module or sub-section
(e.g. "Integration & Hosting", "Module A — Volumes"). Leave Response blank where unknown and tag 🔴.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import json
ACCOUNT = "[Account]"; AREA = "[Product Scope]"; BRAND = "presales-handbook"
# Brand accent: read skills/brand/brands/<BRAND>/brand.json if present, else the
# presales-handbook default navy 112D4E. Users can add their own brand folder.
PLUGIN_ROOT = ""  # SET BY THE SKILL: absolute plugin root (the env var is only a fallback)
_root = PLUGIN_ROOT or os.environ.get("CLAUDE_PLUGIN_ROOT", ".")
try:
    with open(os.path.join(_root, "skills", "brand", "brands", BRAND, "brand.json"), encoding="utf-8") as _fh:
        _b = json.load(_fh)
except (OSError, ValueError):
    _b = {}
ACCENT = ((_b.get("semantic", {}).get("accent") or _b.get("colors", {}).get("navy")
           or _b.get("colors", {}).get("primary") or "112D4E").lstrip("#").upper())

# ROWS = [(group, requirement, description, needed, response, confidence), ...]
ROWS = [
    ("Integration & Hosting", "External system names (ERP, CRM, etc.)", "Names, instances, and info in each", "ROM", "", "🔴"),
    ("Module A — Volumes", "Transactions / records per year", "Drives sizing and licensing", "ROM", "", "🔴"),
    ("Module A — Data", "Master data migration volume", "Migration effort and cleansing", "SOW / CF", "", "🔴"),
]

wb = Workbook(); ws = wb.active; ws.title = "Scoping"
headers = ["Group", "Requirement", "Description", "Needed", "Response", "Confidence"]
widths = [26, 34, 46, 10, 34, 12]

# Title row
ws.merge_cells("A1:F1")
ws["A1"] = f"OSD Scoping — {ACCOUNT} — {AREA}  (ROM = Rough Order of Magnitude · SOW = Statement of Work · CF = Complexity Factor)"
ws["A1"].font = Font(bold=True, color="FFFFFF", size=11)
ws["A1"].fill = PatternFill("solid", fgColor=ACCENT)
ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 28

# Header row
hfill = PatternFill("solid", fgColor=ACCENT)
thin = Side(style="thin", color="DDDDDD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for j, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=2, column=j, value=h)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hfill
    c.alignment = Alignment(vertical="center"); c.border = border
    ws.column_dimensions[chr(64 + j)].width = w

# Data
for i, row in enumerate(ROWS, start=3):
    for j, val in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=val)
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = border
        if headers[j - 1] == "Needed":
            c.font = Font(bold=True)

ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:F{2 + len(ROWS)}"

os.makedirs("output", exist_ok=True)
out = f"output/OSD-Scoping-{ACCOUNT}-{AREA}.xlsx".replace(" ", "_")
wb.save(out)
print(f"Saved: {out}")
