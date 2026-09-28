from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# === BRAND CONFIGURATION ===
# Default brand: presales-handbook (skills/brand/brands/presales-handbook/brand.json).
# If that file is present it is read; otherwise the hardcoded defaults below apply
# (navy 112D4E, yellow FACF39, text 322F2F, Arial). Users can add their own brand
# folder under skills/brand/brands/<name>/ and set BRAND accordingly.
import json

BRAND = "presales-handbook"
# Absolute plugin root — SET BY THE SKILL (discovery-ftd Output B writes the resolved absolute
# path here). It is what makes a custom brand.json and the logo load. ${CLAUDE_PLUGIN_ROOT} is NOT
# expanded inside a Python string, so the env var is only a fallback. If neither is set, the
# presales-handbook default values below apply and a NOTE is printed.
PLUGIN_ROOT = ""
PLUGIN_ROOT = PLUGIN_ROOT or os.environ.get("CLAUDE_PLUGIN_ROOT", "")
if not PLUGIN_ROOT:
    print("NOTE: PLUGIN_ROOT not set — using presales-handbook default brand values, no logo.")
BRAND_DIR = os.path.join(PLUGIN_ROOT or ".", "skills", "brand", "brands", BRAND)

def _hex(value, default):
    return RGBColor.from_string((value or default).lstrip("#").upper())

_brand = {}
try:
    with open(os.path.join(BRAND_DIR, "brand.json"), encoding="utf-8") as fh:
        _brand = json.load(fh)
except (OSError, ValueError):
    pass
_colors   = _brand.get("colors", {})
_semantic = _brand.get("semantic", {})
_fonts    = _brand.get("fonts", {})

ACCENT_COLOR = _hex(_semantic.get("accent") or _colors.get("navy") or _colors.get("primary"), "112D4E")
HIGHLIGHT    = _hex(_semantic.get("highlight") or _colors.get("yellow") or _colors.get("secondary"), "FACF39")
TEXT_COLOR   = _hex(_semantic.get("primary_text") or _colors.get("text"), "322F2F")
HEADING_FONT = _fonts.get("heading", "Arial")
BODY_FONT    = _fonts.get("body", "Arial")
LOGO_PATH    = os.path.join(BRAND_DIR, "assets", "logo-dark.png")  # optional — skipped if absent

doc = Document()

# === PAGE SETUP ===
section = doc.sections[0]
section.page_width  = Cm(21)
section.page_height = Cm(29.7)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)

# === LOGO HEADER ===
header = doc.sections[0].header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = hp.add_run()
if os.path.exists(LOGO_PATH):
    run.add_picture(LOGO_PATH, width=Inches(1.8))

# === TITLE ===
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = title_para.add_run("DISCOVERY QUESTIONNAIRE")
run.font.name  = HEADING_FONT
run.font.size  = Pt(22)
run.font.bold  = True
run.font.color.rgb = ACCENT_COLOR

# === SUBTITLE ===
sub = doc.add_paragraph()
run = sub.add_run("[ACCOUNT NAME] — [PRODUCT SCOPE] Discovery")
run.font.name  = HEADING_FONT
run.font.size  = Pt(13)
run.font.color.rgb = TEXT_COLOR

doc.add_paragraph()

# === MEETING INFO TABLE ===
info_table = doc.add_table(rows=4, cols=2)
info_table.style = "Table Grid"
cells = [
    ("Prospect",     "[Account Name]"),
    ("Date",         "[Meeting Date]"),
    ("Prepared by",  "[SC Name] — [Your Company]"),
    ("Purpose",      "Pre-Discovery Questionnaire — [Product Scope]"),
]
for i, (label, value) in enumerate(cells):
    info_table.rows[i].cells[0].text = label
    info_table.rows[i].cells[1].text = value
    for cell in info_table.rows[i].cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(10)

doc.add_paragraph()

# === INSTRUCTIONS ===
instr = doc.add_paragraph()
run = instr.add_run("Instructions for the prospect")
run.font.name = HEADING_FONT
run.font.size = Pt(11)
run.font.bold = True
run.font.color.rgb = ACCENT_COLOR

doc.add_paragraph(
    "Please complete this questionnaire before our discovery session. "
    "Your answers help us understand your environment and tailor our discussion. "
    "Where we have already noted our understanding, please confirm or correct. "
    "Fields marked * are required."
)

doc.add_paragraph()

# === SECTIONS (replace with generated content) ===
# Each section follows this pattern:
def add_section_heading(doc, title):
    p = doc.add_paragraph()
    run = p.add_run(title.upper())
    run.font.name  = HEADING_FONT
    run.font.size  = Pt(12)
    run.font.bold  = True
    run.font.color.rgb = ACCENT_COLOR
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)

def add_question(doc, question, pre_filled=""):
    q = doc.add_paragraph()
    run = q.add_run(question)
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.font.bold = True
    if pre_filled:
        pf = doc.add_paragraph()
        run2 = pf.add_run(f"Our understanding: {pre_filled}")
        run2.font.name   = BODY_FONT
        run2.font.size   = Pt(9)
        run2.font.italic = True
        run2.font.color.rgb = RGBColor(0x5C, 0x5B, 0x57)
    ans = doc.add_paragraph()
    run3 = ans.add_run("Answer: ")
    run3.font.name = BODY_FONT
    run3.font.size = Pt(10)
    ans.paragraph_format.space_after = Pt(8)

# SECTION CONTENT GOES HERE — replace with skill-generated questions
add_section_heading(doc, "01 — Company & Team Overview")
add_question(doc, "What is your current ERP system and version?", pre_filled="[Pre-filled from research if known]")
add_question(doc, "How many locations / sites are in scope?")
add_question(doc, "Who are the key stakeholders involved in this project?")

# Add more sections here as generated by the skill

# === FOOTER ===
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Confidential — prepared by [Your Company] for [Account Name] | [Date]")
run.font.name = BODY_FONT
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x9A, 0x9A, 0x9A)

# === SAVE ===
os.makedirs("output", exist_ok=True)
doc.save("output/discovery-questionnaire-[account]-[date].docx")
print("Saved: output/discovery-questionnaire-[account]-[date].docx")
